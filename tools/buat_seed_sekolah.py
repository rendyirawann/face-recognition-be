# -*- coding: utf-8 -*-
"""Bangun migrasi seed sekolah Sumut dari berkas XLSX Dinas Pendidikan.

Jalankan ulang bila data sekolahnya diperbarui:

    python tools/buat_seed_sekolah.py "DATA SMA SMK SLB CABDIS DISDIK PROV SU.xlsx"

Keluarannya menimpa api/migrations/0020_seed_sekolah_sumut.sql.
"""
import io
import os
import re
import sys
import unicodedata
from collections import OrderedDict

from openpyxl import load_workbook

KELUARAN = os.path.join("api", "migrations", "0020_seed_sekolah_sumut.sql")

# bcrypt cost 12, format $2y$ agar cocok dengan Hash::check() Laravel.
# Dibangkitkan: php -r "password_hash(..., PASSWORD_BCRYPT, ['cost'=>12])"
HASH_CAPDIS = "$2y$12$u/YK3TjLn.OoFRb2K/0QreT4ArTOPhZwJTlczLyt5aIUuE5H5oeNa"  # Capdis#2026
HASH_DISDIK = "$2y$12$XusURJP3EkaOE8nT6z2lmewrbJs9o6xfN6OU1OrA2f0U36bKqqhF."  # Disdik#2026

# Romawi -> angka, untuk kode wilayah yang bisa diurutkan.
ROMAWI = {
    "I": 1, "II": 2, "III": 3, "IV": 4, "V": 5, "VI": 6, "VII": 7,
    "VIII": 8, "IX": 9, "X": 10, "XI": 11, "XII": 12, "XIII": 13, "XIV": 14,
}


def slug(teks: str) -> str:
    t = unicodedata.normalize("NFKD", teks).encode("ascii", "ignore").decode()
    t = re.sub(r"[^a-zA-Z0-9]+", "-", t).strip("-").lower()
    return re.sub(r"-{2,}", "-", t)


def sql(teks) -> str:
    if teks is None:
        return "NULL"
    return "'" + str(teks).replace("'", "''") + "'"


def normalkan_jenjang(nilai: str) -> str:
    """SMKS dan 'SMK NEGERI 1 LAHEWA' sama-sama berarti SMK.

    Dua baris di sumber datanya salah isi — kolom jenjang tertulis nama
    sekolah. Prefiksnya tetap benar, jadi itu yang dipakai.
    """
    v = (nilai or "").strip().upper()
    for j in ("SMK", "SMA", "SLB"):
        if v.startswith(j):
            return j
    raise ValueError(f"jenjang tidak dikenal: {nilai!r}")


def main(sumber: str) -> int:
    wb = load_workbook(sumber, read_only=True)

    wilayah = OrderedDict()   # nama sheet -> {kode, nomor, kabupaten: set}
    sekolah = []             # (npsn, nama, slug, jenjang, status, kab, guru, siswa)
    kabupaten = OrderedDict()  # nama kab -> sheet wilayah
    kode_kab = {}              # nama kab -> kode pendek
    dipakai_slug = set()
    urut = 0

    for sheet in wb.sheetnames:
        m = re.search(r"WILAYAH\s+([IVX]+)", sheet.upper())
        if not m:
            print(f"lewati sheet tanpa nomor wilayah: {sheet}")
            continue
        nomor = ROMAWI[m.group(1)]
        wilayah[sheet] = {"nomor": nomor, "kode": f"CABDIS{nomor:02d}"}

        for i, row in enumerate(wb[sheet].iter_rows(values_only=True)):
            if i == 0:
                continue
            nama, jen, sts, kab, guru, siswa = (tuple(row) + (None,) * 6)[:6]
            if not nama or not str(nama).strip():
                continue

            nama = str(nama).strip()
            kab = str(kab).strip() if kab else None
            if kab and kab not in kabupaten:
                kabupaten[kab] = sheet
                # Berurut, BUKAN potongan slug: `regions.code` hanya 10
                # karakter sedangkan seluruh nama berawalan "Kabupaten "
                # atau "Kota ", sehingga potongannya identik dan bentrok.
                kode_kab[kab] = f"K{len(kode_kab) + 1:02d}"


            urut += 1
            # NPSN sintetis: berawalan huruf sehingga TIDAK MUNGKIN bentrok
            # dengan NPSN asli yang seluruhnya angka. Lihat catatan di kepala
            # berkas SQL keluaran.
            npsn = f"SMT{urut:07d}"

            s = slug(nama)[:200] or f"sekolah-{urut}"
            asal = s
            n = 2
            while s in dipakai_slug:
                s = f"{asal}-{n}"
                n += 1
            dipakai_slug.add(s)

            sekolah.append((
                npsn, nama, s, normalkan_jenjang(jen),
                (str(sts).strip().lower() if sts else "negeri"),
                kab,
                int(guru) if guru and str(guru).isdigit() else 0,
                int(siswa) if siswa and str(siswa).isdigit() else 0,
            ))

    b = []
    w = b.append
    w("-- =====================================================================")
    w("-- 0020 : Seed wilayah Cabang Dinas, sekolah, dan akun pemantau")
    w("--")
    w(f"-- Dibangkitkan tools/buat_seed_sekolah.py dari {os.path.basename(sumber)}")
    w(f"-- {len(sekolah)} sekolah, {len(kabupaten)} kabupaten/kota, {len(wilayah)} wilayah.")
    w("--")
    w("-- NPSN BELUM ASLI. Sumber datanya tidak memuat kolom NPSN, sedangkan")
    w("-- kolomnya NOT NULL UNIQUE. Nilai sementara berbentuk SMT0000001 —")
    w("-- berawalan huruf sehingga tidak mungkin bentrok dengan NPSN asli yang")
    w("-- seluruhnya angka, dan mudah dicari saat penggantian:")
    w("--")
    w("--   SELECT count(*) FROM schools WHERE npsn LIKE 'SMT%';")
    w("--")
    w("-- Ganti dengan NPSN Kemdikbud sebelum sistem dipakai sungguhan;")
    w("-- tanpa itu, impor data resmi akan menghasilkan sekolah ganda.")
    w("-- =====================================================================")
    w("")
    w("-- Provinsi")
    w("INSERT INTO regions (code, name, kind) VALUES ('12', 'Sumatera Utara', 'provinsi')")
    w("ON CONFLICT (code) DO NOTHING;")
    w("")
    w("-- Cabang Dinas, anak provinsi")
    w("INSERT INTO regions (code, name, kind, parent_id) VALUES")
    baris = []
    for sheet, info in wilayah.items():
        romawi = [k for k, v in ROMAWI.items() if v == info["nomor"]][0]
        nama_wil = f"Cabang Dinas Wilayah {romawi}"
        baris.append(
            f"    ({sql(info['kode'])}, {sql(nama_wil)}, 'cabang_dinas', "
            f"(SELECT id FROM regions WHERE code = '12'))"
        )
    w(",\n".join(baris))
    w("ON CONFLICT (code) DO NOTHING;")
    w("")
    w("-- Kabupaten/kota, anak Cabang Dinas masing-masing")
    w("INSERT INTO regions (code, name, kind, parent_id) VALUES")
    baris = []
    for kab, sheet in kabupaten.items():
        kind = "kota" if kab.lower().startswith("kota") else "kabupaten"
        kode = kode_kab[kab]
        baris.append(
            f"    ({sql(kode)}, {sql(kab)}, {sql(kind)}, "
            f"(SELECT id FROM regions WHERE code = {sql(wilayah[sheet]['kode'])}))"
        )
    w(",\n".join(baris))
    w("ON CONFLICT (code) DO NOTHING;")
    w("")
    w(f"-- {len(sekolah)} sekolah")
    w("INSERT INTO schools (npsn, name, slug, jenjang, status, region_id, settings) VALUES")
    baris = []
    for npsn, nama, s, jen, sts, kab, guru, siswa in sekolah:
        wil = f"(SELECT id FROM regions WHERE code = {sql(kode_kab[kab])})" if kab else "NULL"
        # Jumlah guru & siswa dari sumber disimpan di settings: dipakai
        # membangkitkan data dummy yang proporsional, bukan angka karangan.
        cfg = '{"guru": %d, "siswa": %d}' % (guru, siswa)
        baris.append(
            f"    ({sql(npsn)}, {sql(nama)}, {sql(s)}, {sql(jen)}, {sql(sts)}, "
            f"{wil}, {sql(cfg)}::jsonb)"
        )
    w(",\n".join(baris))
    w("ON CONFLICT (npsn) DO NOTHING;")
    w("")

    # ----------------------------------------------------------------
    # Akun pemantau
    # ----------------------------------------------------------------
    w("-- Akun Cabang Dinas: satu per wilayah. Cakupan sekolahnya diturunkan")
    w("-- dari users.region_id lewat hierarki regions, bukan daftar manual.")
    w("--")
    w("-- Kata sandi awal sama untuk keempat belas akun dan wajib diganti saat")
    w("-- login pertama. Sandi berbeda per akun tidak menambah keamanan selama")
    w("-- semuanya diserahkan lewat kanal yang sama.")
    w("--")
    w("--   pengguna : capdis01 .. capdis14")
    w("--   sandi    : Capdis#2026")
    w("INSERT INTO users (name, username, email, password, region_id,")
    w("                   is_active, email_verified_at, must_change_password,")
    w("                   created_at, updated_at) VALUES")
    baris = []
    for sheet, info in wilayah.items():
        romawi = [k for k, v in ROMAWI.items() if v == info["nomor"]][0]
        akun = f"capdis{info['nomor']:02d}"
        baris.append(
            f"    ({sql('Cabang Dinas Wilayah ' + romawi)}, {sql(akun)}, "
            f"{sql(akun + '@disdik.sumutprov.go.id')}, {sql(HASH_CAPDIS)}, "
            f"(SELECT id FROM regions WHERE code = {sql(info['kode'])}), "
            f"TRUE, NOW(), TRUE, NOW(), NOW())"
        )
    w(",\n".join(baris))
    w("ON CONFLICT (username) DO NOTHING;")
    w("")
    w("INSERT INTO model_has_roles (role_id, model_type, model_id)")
    w("SELECT r.id, 'App\\Models\\User', u.id")
    w("  FROM roles r, users u")
    w(" WHERE r.name = 'capdis' AND r.guard_name = 'web'")
    w("   AND u.username LIKE 'capdis__'")
    w("ON CONFLICT DO NOTHING;")
    w("")
    w("-- Akun Dinas Pendidikan provinsi: memantau seluruh wilayah.")
    w("-- region_id NULL — cakupannya diturunkan dari peran, bukan satu wilayah.")
    w("--")
    w("--   pengguna : disdik")
    w("--   sandi    : Disdik#2026")
    w("INSERT INTO users (name, username, email, password, is_active,")
    w("                   email_verified_at, must_change_password,")
    w("                   created_at, updated_at) VALUES")
    w("    ('Dinas Pendidikan Provinsi Sumatera Utara', 'disdik',")
    w(f"     'disdik@disdik.sumutprov.go.id', {sql(HASH_DISDIK)},")
    w("     TRUE, NOW(), TRUE, NOW(), NOW())")
    w("ON CONFLICT (username) DO NOTHING;")
    w("")
    w("INSERT INTO model_has_roles (role_id, model_type, model_id)")
    w("SELECT r.id, 'App\\Models\\User', u.id")
    w("  FROM roles r, users u")
    w(" WHERE r.name = 'admin_dinas' AND r.guard_name = 'web'")
    w("   AND u.username = 'disdik'")
    w("ON CONFLICT DO NOTHING;")
    w("")

    io.open(KELUARAN, "w", encoding="utf-8", newline="").write("\n".join(b) + "\n")
    print(f"{KELUARAN}: {len(sekolah)} sekolah, {len(kabupaten)} kab/kota, "
          f"{len(wilayah)} wilayah")
    return 0


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(2)
    sys.exit(main(sys.argv[1]))
